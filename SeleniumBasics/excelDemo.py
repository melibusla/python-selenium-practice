import openpyxl
book = openpyxl.load_workbook('/home/melina/Documents/testdata.xlsx')

sheet = book.active
#empty dictionary
data = {}

#get the number of rows and columns
print(sheet.max_row)
print(sheet.max_column)

cell = sheet.cell(row=1 ,column=2)
print(cell.value)

sheet.cell(row=2 ,column=2).value = "Mel"
print(sheet.cell(row=2 ,column=2).value)

print(sheet['A5'].value)

#Print the values of the complete table
# for i in range(1, sheet.max_row + 1):
#     for j in range(1, sheet.max_column + 1):
#         print(sheet.cell(row=i, column=j).value)

#Print the values of a table with a condition
for i in range(1, sheet.max_row + 1): #rows
    if sheet.cell(row=i, column=1).value == "testCase2":
        # for j in range(1, sheet.max_column + 1): #columns
        for j in range(2, sheet.max_column + 1):  # columns with no column A
            # print(sheet.cell(row=i, column=j).value)
            #Dic["lastname"]=B
            data[sheet.cell(row=1, column=j).value] = sheet.cell(
                row=i,
                column=j
            ).value
        print(data)