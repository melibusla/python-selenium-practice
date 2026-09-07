import os
import time

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.wait import WebDriverWait
import openpyxl


def update_excel_data(file_path, searchTerm, colName, new_value):
    book = openpyxl.load_workbook(file_path)
    sheet = book.active

    column = None
    for i in range(1, sheet.max_column + 1):
        header = sheet.cell(row=1, column=i).value
        if str(header).strip().casefold() == colName.strip().casefold():
            column = i
            break

    row = None
    for i in range(1, sheet.max_row + 1):
        for j in range(1, sheet.max_column + 1):
            if sheet.cell(row=i, column=j).value == searchTerm:
                row = i
                break
        if row is not None:
            break

    if column is None:
        headers = [sheet.cell(row=1, column=i).value
                   for i in range(1, sheet.max_column + 1)]
        raise ValueError(
            f"Column {colName!r} was not found in {file_path}. "
            f"Available headers: {headers!r}"
        )
    if row is None:
        raise ValueError(
            f"Search value {searchTerm!r} was not found in {file_path}."
        )

    sheet.cell(row=row, column=column).value = new_value
    book.save(file_path)


driver = webdriver.Chrome()
driver.implicitly_wait(5)
driver.get("https://rahulshettyacademy.com/upload-download-test/index.html")
file_path = "/home/melina/Downloads/download.xlsx"
fruit_name = "Apple"
new_value = "999"

# Download the file
if os.path.exists(file_path):
    os.remove(file_path)
driver.find_element(By.ID, "downloadButton").click()

# Wait until the download is complete
deadline = time.time() + 10
while not os.path.exists(file_path) and time.time() < deadline:
    time.sleep(0.1)
if not os.path.exists(file_path):
    raise TimeoutError(f"Download did not complete: {file_path}")

# Edit the file
update_excel_data(file_path, fruit_name, "Price", new_value)

# Upload the file
file_input = driver.find_element(By.CSS_SELECTOR, "input[type='file']")
file_input.send_keys(file_path)

# Catch the toast message
wait = WebDriverWait(driver, 5)
toast_locator = (By.CSS_SELECTOR, ".Toastify__toast-body div:nth-child(2)")
wait.until(expected_conditions.visibility_of_element_located(toast_locator))
print(driver.find_element(*toast_locator).text)

# Retrieve price of a row
price_column = driver.find_element(By.XPATH, "//div[text()='Price']").get_attribute("data-column-id")
actual_price = driver.find_element(
    By.XPATH,
    "//div[text()='" + fruit_name + "']/parent::div/parent::div/"
    "div[@id='cell-" + price_column + "-undefined']"
).text
assert actual_price == new_value
