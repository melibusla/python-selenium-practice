import os
import time

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.wait import WebDriverWait
import openpyxl

# Isolated method
def update_excel_data(file_path, searchTerm, colName, new_value):
    book = openpyxl.load_workbook(file_path)
    sheet = book.active
    dataDict = {}

# Obtain the column number and add it to a dictionary
    for i in range(1, sheet.max_column + 1):
        # colName fix: match headers regardless of case or extra spaces
        if str(sheet.cell(row=1, column=i).value).strip().casefold() == colName.strip().casefold():
            dataDict["col"] = i
# Obtain the row
    for i in range(1, sheet.max_row + 1):
        for j in range(1, sheet.max_column + 1):
            if sheet.cell(row=i, column=j).value == searchTerm:
                dataDict["row"] = i
# Edit the file
    sheet.cell(row=dataDict["row"], column=dataDict["col"]).value = new_value
# save
    book.save(file_path)

driver = webdriver.Chrome()
driver.implicitly_wait(5)
driver.get("https://rahulshettyacademy.com/upload-download-test/index.html")
file_path = "/home/melina/Downloads/download.xlsx"
fruit_name = "Apple"
new_value = "999"

#Download the file
if os.path.exists(file_path):
    os.remove(file_path) #Deletes the old file
driver.find_element(By.ID, "downloadButton").click()

# Wait for the download to finish
download_timeout = time.time() + 10
while not os.path.exists(file_path) and time.time() < download_timeout:
    time.sleep(0.1)
if not os.path.exists(file_path):
    raise FileNotFoundError("The Excel download did not finish: " + file_path)

#Edit the file
update_excel_data(file_path, fruit_name, "Price", new_value)

#Upload the file
file_input = driver.find_element(By.CSS_SELECTOR, "input[type='file']")
file_input.send_keys(file_path)

#catch the toast message
wait = WebDriverWait(driver, 5)
toast_locator = (By.CSS_SELECTOR, ".Toastify__toast-body div:nth-child(2)")
wait.until(expected_conditions.visibility_of_element_located(toast_locator))
print(driver.find_element(*toast_locator).text)

#Retrieve price of a row
#Retrieve price column attribute
price_column = driver.find_element(By.XPATH, "//div[text()='Price']").get_attribute("data-column-id")
actual_price = driver.find_element(By.XPATH, "//div[text()='"+fruit_name+"']/parent::div/parent::div/div[@id='cell-"+price_column+"-undefined']").text
assert actual_price == new_value